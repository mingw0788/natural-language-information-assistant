from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass
from datetime import UTC, datetime
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from info_assistant.domain.models import TaskSpec
from info_assistant.storage.database import Database

FORMATS = {"text", "csv", "xlsx", "json", "markdown"}
HEADERS = [
    "task_id",
    "task_version_id",
    "run_id",
    "generated_at",
    "view",
    "entity_id",
    "entity",
    "field",
    "value",
    "status",
    "origin",
    "unit",
    "period",
    "qualifiers",
    "source_urls",
    "source_quotes",
    "conditions_and_limits",
    "record_type",
    "title",
    "url",
    "excerpt",
    "fetched_at",
]


@dataclass(frozen=True)
class ExportArtifact:
    export_id: str
    filename: str
    content_type: str
    content: bytes
    record_count: int


class ExportService:
    def __init__(self, database: Database) -> None:
        self.database = database

    def export(
        self,
        *,
        task_id: str,
        task_version_id: str,
        run_id: str | None,
        spec: TaskSpec,
        records: list[dict],
        format: str,
        view: str = "all",
        selected_entity_ids: list[str] | None = None,
        selected_fields: list[str] | None = None,
        sources: list[dict] | None = None,
    ) -> ExportArtifact:
        file_format = format.strip().lower()
        if file_format not in FORMATS:
            raise ValueError("导出格式仅支持 text、csv、xlsx、json 和 markdown")
        selected_entities = set(selected_entity_ids) if selected_entity_ids is not None else None
        selected_field_set = {value.casefold() for value in selected_fields or []}
        filtered = []
        for record in records:
            if selected_entities is not None and record.get("entity_id") not in selected_entities:
                continue
            selected_observations = [
                item
                for item in record.get("observations", [])
                if not selected_field_set or item.get("field", "").casefold() in selected_field_set
            ]
            if selected_field_set and not selected_observations:
                continue
            filtered.append({**record, "observations": selected_observations})
        source_records = []
        if not records and view == "all" and selected_entities is None and not selected_field_set:
            source_records = [
                {
                    "record_type": "source",
                    **{
                        key: str(source.get(key) or "")
                        for key in ("title", "url", "excerpt", "fetched_at")
                    },
                }
                for source in sources or []
            ]
        generated_at = datetime.now(UTC).isoformat()
        metadata = {
            "task_id": task_id,
            "task_version_id": task_version_id,
            "run_id": run_id,
            "generated_at": generated_at,
            "view": view,
            "task": spec.raw_input,
            "summary": spec.summary,
            "fields": [field.model_dump(mode="json") for field in spec.fields],
            "required_conditions": [
                item.model_dump(mode="json") for item in spec.required_conditions
            ],
            "preferences": [item.model_dump(mode="json") for item in spec.preferences],
            "assumptions": spec.assumptions,
            "source_constraints": spec.source_constraints,
            "time_window": spec.time_window.model_dump(mode="json") if spec.time_window else None,
            "stop_conditions": spec.stop_conditions,
            "display_limit": spec.display_limit,
            "selected_entity_ids": (
                sorted(selected_entities) if selected_entities is not None else None
            ),
            "selected_fields": sorted(selected_field_set) if selected_field_set else None,
            "source_count": len(source_records),
            "limitations": [
                "搜索摘要只作为线索；导出的来源摘录来自已保存的来源快照。",
                "结果可能部分完成；未知字段不会自动当作不存在。",
                "CSV/XLSX 中以公式字符开头的文本已作为纯文本处理。",
            ],
        }
        rows = self._rows(
            task_id,
            task_version_id,
            run_id,
            generated_at,
            view,
            filtered,
            spec,
            selected_field_set,
        )
        if source_records:
            metadata["limitations"].append("来源文档尚未作结构化核实，不计入正式结果数量。")
            rows.extend(
                {
                    "task_id": task_id,
                    "task_version_id": task_version_id,
                    "run_id": run_id or "",
                    "generated_at": generated_at,
                    "view": view,
                    "status": "source_only",
                    "conditions_and_limits": "来源文档尚未作结构化核实，不计入正式结果数量。",
                    **source,
                }
                for source in source_records
            )
        metadata["record_count"] = len(rows)
        if file_format == "csv":
            content = self._csv(rows)
            content_type = "text/csv; charset=utf-8-sig"
            extension = "csv"
        elif file_format == "xlsx":
            content = self._xlsx(rows, metadata)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            extension = "xlsx"
        elif file_format == "json":
            content = json.dumps(
                {"metadata": metadata, "records": filtered, "sources": source_records},
                ensure_ascii=False, indent=2,
            ).encode("utf-8")
            content_type = "application/json; charset=utf-8"
            extension = "json"
        elif file_format == "markdown":
            content = self._markdown(metadata, rows).encode("utf-8")
            content_type = "text/markdown; charset=utf-8"
            extension = "md"
        else:
            content = self._text(metadata, rows).encode("utf-8")
            content_type = "text/plain; charset=utf-8"
            extension = "txt"

        export_id = str(uuid4())
        selection = {
            "view": view,
            "selected_entity_ids": metadata["selected_entity_ids"],
            "selected_fields": metadata["selected_fields"],
            "sort": "entity, field",
        }
        with self.database.transaction() as connection:
            connection.execute(
                "INSERT INTO exports(id, task_version_id, selection_json, format, generated_at, "
                "file_hash, record_count) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    export_id,
                    task_version_id,
                    json.dumps(selection, ensure_ascii=False),
                    file_format,
                    generated_at,
                    hashlib.sha256(content).hexdigest(),
                    len(rows),
                ),
            )
        return ExportArtifact(
            export_id=export_id,
            filename=f"information-assistant-{task_id[:8]}.{extension}",
            content_type=content_type,
            content=content,
            record_count=len(rows),
        )

    @staticmethod
    def _rows(
        task_id: str,
        task_version_id: str,
        run_id: str | None,
        generated_at: str,
        view: str,
        records: list[dict],
        spec: TaskSpec,
        selected_fields: set[str],
    ) -> list[dict[str, str]]:
        conditions = [item.text for item in spec.required_conditions]
        conditions.extend(spec.assumptions)
        conditions.extend(spec.stop_conditions)
        conditions.extend(spec.source_constraints)
        if spec.time_window:
            window_label = spec.time_window.label or (
                f"{spec.time_window.start.isoformat()} 至 {spec.time_window.end.isoformat()}"
            )
            conditions.append(f"时间范围：{window_label}")
        rows = []
        for record in records:
            for observation in record.get("observations", []):
                if (
                    selected_fields
                    and observation.get("field", "").casefold() not in selected_fields
                ):
                    continue
                evidence = observation.get("evidence", [])
                row = {
                    "task_id": task_id,
                    "task_version_id": task_version_id,
                    "run_id": run_id or "",
                    "generated_at": generated_at,
                    "view": view,
                    "entity_id": record.get("entity_id", ""),
                    "entity": record.get("entity", ""),
                    "field": observation.get("field", ""),
                    "value": observation.get("value", ""),
                    "status": observation.get("status", record.get("status", "")),
                    "origin": observation.get("origin", ""),
                    "unit": observation.get("unit", "") or "",
                    "period": observation.get("period", "") or "",
                    "qualifiers": "；".join(observation.get("qualifiers", [])),
                    "source_urls": "；".join(item.get("url", "") or "" for item in evidence),
                    "source_quotes": "\n---\n".join(item.get("quote", "") for item in evidence),
                    "conditions_and_limits": "；".join(conditions),
                    "record_type": "observation",
                }
                rows.append(row)
        return rows

    @staticmethod
    def _safe_cell(value: object) -> str:
        text = "" if value is None else str(value)
        first = text.lstrip(" \t\r\n")[:1]
        if first in {"=", "+", "-", "@", "\t", "\r"}:
            return "'" + text
        return text

    @classmethod
    def _csv(cls, rows: list[dict[str, str]]) -> bytes:
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=HEADERS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: cls._safe_cell(row.get(key)) for key in HEADERS})
        return output.getvalue().encode("utf-8-sig")

    @classmethod
    def _xlsx(cls, rows: list[dict[str, str]], metadata: dict) -> bytes:
        workbook = Workbook()
        results = workbook.active
        results.title = "结果"
        results.append(HEADERS)
        meta = workbook.create_sheet("元信息")
        meta.append(["字段", "值"])
        meta_keys = [
            "task_version_id",
            "task_id",
            "run_id",
            "generated_at",
            "view",
            "task",
            "summary",
            "required_conditions",
            "assumptions",
            "source_constraints",
            "time_window",
            "stop_conditions",
            "limitations",
            "source_count",
        ]
        for key in meta_keys:
            value = metadata.get(key)
            serialized = (
                json.dumps(value, ensure_ascii=False)
                if isinstance(value, (list, dict))
                else value or ""
            )
            meta.append([cls._safe_cell(key), cls._safe_cell(serialized)])
        for row in rows:
            results.append([cls._safe_cell(row.get(header, "")) for header in HEADERS])
        header_fill = PatternFill("solid", fgColor="DCE7F7")
        for sheet in (results, meta):
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="172033")
                cell.fill = header_fill
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cell.number_format = "@"
                        cell.data_type = "s"
            for column in sheet.columns:
                letter = column[0].column_letter
                max_width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
                sheet.column_dimensions[letter].width = max(14, max_width)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _markdown(metadata: dict, rows: list[dict[str, str]]) -> str:
        def escape(value: object) -> str:
            return str(value or "").replace("|", "\\|").replace("\r", " ").replace("\n", "<br>")

        lines = [
            f"# {metadata.get('summary') or metadata.get('task') or '信息采集结果'}",
            "",
            f"- 任务：`{metadata['task_id']}`",
            f"- 任务版本：`{metadata['task_version_id']}`",
            f"- 运行：`{metadata.get('run_id') or '无'}`",
            f"- 生成时间：{metadata['generated_at']}",
            f"- 当前视图：{metadata['view']}",
            "- 条件与限制："
            + (
                "；".join(
                    metadata["assumptions"]
                    + metadata["stop_conditions"]
                    + metadata["source_constraints"]
                )
                or "无额外说明"
            ),
            "",
        ]
        if metadata["source_count"]:
            lines.extend([
                "## 来源文档（record_type=source）",
                "",
                "来源文档尚未作结构化核实，不计入正式结果数量。",
                "",
                "| 标题 | 来源链接 | 原文摘录 | 读取时间 |",
                "| --- | --- | --- | --- |",
            ])
        else:
            lines.extend([
                "| 对象 | 字段 | 值 | 状态 | 来源摘录 | 来源链接 |",
                "| --- | --- | --- | --- | --- | --- |",
            ])
        for row in rows:
            keys = (
                ("title", "url", "excerpt", "fetched_at")
                if row["record_type"] == "source"
                else ("entity", "field", "value", "status", "source_quotes", "source_urls")
            )
            lines.append(
                "| "
                + " | ".join(escape(row[key]) for key in keys)
                + " |"
            )
        lines.extend(["", "搜索摘要只作为线索；结果可能部分完成，未知字段不会自动当作不存在。"])
        return "\n".join(lines) + "\n"

    @staticmethod
    def _text(metadata: dict, rows: list[dict[str, str]]) -> str:
        lines = [
            str(metadata.get("summary") or metadata.get("task") or "信息采集结果"),
            f"任务版本：{metadata['task_version_id']}",
            f"运行：{metadata.get('run_id') or '无'}",
            f"生成时间：{metadata['generated_at']}",
            f"视图：{metadata['view']}",
            "",
        ]
        for row in rows:
            if row["record_type"] == "source":
                lines.extend([
                    "来源文档（record_type=source，尚未作结构化核实）",
                    f"标题：{row['title']}",
                    f"来源：{row['url']}",
                    f"读取时间：{row['fetched_at']}",
                    f"原文摘录：{row['excerpt']}",
                    "",
                ])
                continue
            lines.append(f"{row['entity']} · {row['field']}：{row['value']} [{row['status']}]")
            if row["source_quotes"]:
                lines.append(f"来源摘录：{row['source_quotes']}")
            if row["source_urls"]:
                lines.append(f"来源：{row['source_urls']}")
        lines.extend(["", "搜索摘要只作为线索；结果可能部分完成，未知字段不会自动当作不存在。"])
        return "\n".join(lines)
