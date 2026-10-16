from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from info_assistant.domain.models import TaskField, TaskSpec
from info_assistant.services.export_service import ExportService
from info_assistant.storage.database import Database
from info_assistant.storage.repository import Repository


def make_export_service(tmp_path: Path) -> tuple[ExportService, str, str]:
    database = Database(tmp_path / "export.sqlite3")
    database.migrate()
    task = Repository(database).create_task(
        TaskSpec(
            raw_input="查产品价格",
            fields=[TaskField(name="价格")],
            required_conditions=[],
            assumptions=["未指定地区，不推断地区"],
        )
    )
    return ExportService(database), task.id, task.latest_version.id


def sample_records() -> list[dict]:
    return [
        {
            "entity_id": "90071992547409931234",
            "entity": "=HYPERLINK(\"https://evil.example\")",
            "status": "confirmed",
            "reasons": ["有精确来源摘录"],
            "observations": [
                {
                    "field": "价格",
                    "value": "+cmd|' /C calc'!A0",
                    "status": "confirmed",
                    "unit": None,
                    "period": "每月",
                    "origin": "source",
                    "evidence": [
                        {
                            "quote": "价格：+cmd|' /C calc'!A0",
                            "url": "https://example.test/pricing",
                            "title": "Example pricing",
                        }
                    ],
                }
            ],
        }
    ]


def test_csv_escapes_formula_text_and_includes_trace_metadata(tmp_path: Path) -> None:
    service, task_id, version_id = make_export_service(tmp_path)

    artifact = service.export(
        task_id=task_id,
        task_version_id=version_id,
        run_id="run-123",
        spec=TaskSpec(raw_input="查产品价格", assumptions=["未指定地区，不推断地区"]),
        records=sample_records(),
        format="csv",
        view="confirmed",
    )

    content = artifact.content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(content)))
    assert rows[0]["entity"].startswith("'=")
    assert rows[0]["value"].startswith("'+")
    assert rows[0]["task_version_id"] == version_id
    assert rows[0]["view"] == "confirmed"
    assert "未指定地区，不推断地区" in rows[0]["conditions_and_limits"]


def test_xlsx_uses_text_cells_and_formula_safety(tmp_path: Path) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    artifact = service.export(
        task_id=task_id,
        task_version_id=version_id,
        run_id=None,
        spec=TaskSpec(raw_input="查产品价格"),
        records=sample_records(),
        format="xlsx",
    )

    workbook = load_workbook(io.BytesIO(artifact.content), data_only=False)
    cell = workbook["结果"]["A2"]
    value_index = [cell.value for cell in workbook["结果"][1]].index("entity") + 1
    entity = workbook["结果"].cell(row=2, column=value_index)
    assert cell.data_type == "s"
    assert str(workbook["结果"].cell(row=2, column=1).value) == task_id
    assert entity.data_type == "s"
    assert entity.value.startswith("'=")
    assert workbook["元信息"]["B2"].value == version_id


def test_json_and_markdown_are_auditable_and_export_is_recorded(tmp_path: Path) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    kwargs = {
        "task_id": task_id,
        "task_version_id": version_id,
        "run_id": "run-xyz",
        "spec": TaskSpec(raw_input="查产品价格", assumptions=["时间未指定"]),
        "records": sample_records(),
        "format": "json",
        "view": "all",
    }

    artifact = service.export(**kwargs)
    payload = json.loads(artifact.content)
    assert payload["metadata"]["task_id"] == task_id
    assert payload["metadata"]["run_id"] == "run-xyz"
    assert payload["metadata"]["assumptions"] == ["时间未指定"]
    assert payload["records"][0]["observations"][0]["evidence"][0]["url"]

    markdown = service.export(**{**kwargs, "format": "markdown"}).content.decode("utf-8")
    assert "任务版本" in markdown
    assert "来源摘录" in markdown
    assert "\|" in markdown
    with service.database.connection() as connection:
        count = connection.execute("SELECT COUNT(*) FROM exports").fetchone()[0]
    assert count == 2


def sample_sources() -> list[dict]:
    return [{
        "snapshot_id": "snapshot-1",
        "title": "Python 官方文档",
        "url": "https://docs.python.org/zh-cn/3/",
        "excerpt": "这里是已读取的原文，尚未核实为结构化事实。",
        "fetched_at": "2026-10-15T12:00:00Z",
    }]


@pytest.mark.parametrize("file_format", ["json", "markdown", "text", "csv", "xlsx"])
def test_source_only_collection_exports_documents_without_inventing_entities(
    tmp_path: Path, file_format: str
) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    artifact = service.export(
        task_id=task_id,
        task_version_id=version_id,
        run_id="source-run",
        spec=TaskSpec(raw_input="找 Python 官方文档"),
        records=[],
        sources=sample_sources(),
        format=file_format,
    )

    assert artifact.record_count == 1
    if file_format == "json":
        payload = json.loads(artifact.content)
        assert payload["records"] == []
        assert payload["sources"] == [{
            "record_type": "source",
            "title": "Python 官方文档",
            "url": "https://docs.python.org/zh-cn/3/",
            "excerpt": "这里是已读取的原文，尚未核实为结构化事实。",
            "fetched_at": "2026-10-15T12:00:00Z",
        }]
        assert payload["metadata"]["source_count"] == 1
    elif file_format in {"csv", "xlsx"}:
        if file_format == "csv":
            rows = list(csv.DictReader(io.StringIO(artifact.content.decode("utf-8-sig"))))
        else:
            sheet = load_workbook(io.BytesIO(artifact.content))["结果"]
            values = list(sheet.values)
            rows = [dict(zip(values[0], row, strict=True)) for row in values[1:]]
        assert len(rows) == 1
        assert rows[0]["record_type"] == "source"
        assert rows[0]["title"] == "Python 官方文档"
        assert rows[0]["url"] == "https://docs.python.org/zh-cn/3/"
        assert rows[0]["excerpt"] == "这里是已读取的原文，尚未核实为结构化事实。"
        assert rows[0]["fetched_at"] == "2026-10-15T12:00:00Z"
        assert not rows[0]["entity_id"]
        assert not rows[0]["value"]
        assert rows[0]["status"] != "confirmed"
    else:
        text = artifact.content.decode("utf-8")
        assert "来源文档" in text
        assert "record_type=source" in text
        for key in ("title", "url", "excerpt", "fetched_at"):
            assert sample_sources()[0][key] in text
    with service.database.connection() as connection:
        row_count = connection.execute(
            "SELECT record_count FROM exports WHERE id=?", (artifact.export_id,)
        ).fetchone()[0]
    assert row_count == 1


@pytest.mark.parametrize("selection", [
    {"view": "confirmed"},
    {"view": "pending_review"},
    {"selected_entity_ids": []},
    {"selected_entity_ids": ["not-selected"]},
    {"selected_fields": ["价格"]},
])
def test_source_fallback_does_not_override_entity_view_or_field_selection(
    tmp_path: Path, selection: dict
) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    artifact = service.export(
        task_id=task_id, task_version_id=version_id, run_id="source-run",
        spec=TaskSpec(raw_input="查价格"), records=[], sources=sample_sources(),
        format="json", **selection,
    )
    payload = json.loads(artifact.content)
    assert payload["sources"] == []
    assert payload["metadata"]["source_count"] == 0
    assert artifact.record_count == 0


def test_source_fallback_does_not_mix_documents_into_structured_export(tmp_path: Path) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    artifact = service.export(
        task_id=task_id, task_version_id=version_id, run_id="source-run",
        spec=TaskSpec(raw_input="查价格"), records=sample_records(), sources=sample_sources(),
        format="json",
    )
    payload = json.loads(artifact.content)
    assert payload["sources"] == []
    assert payload["records"] == sample_records()
    assert artifact.record_count == 1


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_source_document_spreadsheet_cells_escape_formula_text(
    tmp_path: Path, file_format: str
) -> None:
    service, task_id, version_id = make_export_service(tmp_path)
    sources = [{**sample_sources()[0], "title": "=1+1", "excerpt": "+cmd|' /C calc'!A0"}]
    artifact = service.export(
        task_id=task_id, task_version_id=version_id, run_id="source-run",
        spec=TaskSpec(raw_input="读取资料"), records=[], sources=sources, format=file_format,
    )
    if file_format == "csv":
        row = next(csv.DictReader(io.StringIO(artifact.content.decode("utf-8-sig"))))
    else:
        sheet = load_workbook(io.BytesIO(artifact.content), data_only=False)["结果"]
        row = dict(zip(next(sheet.values), list(sheet.values)[1], strict=True))
        assert all(cell.data_type != "f" for cell in sheet[2])
    assert row["title"] == "'=1+1"
    assert row["excerpt"] == "'+cmd|' /C calc'!A0"
